import json
from pathlib import Path

import openpyxl

source = Path(r"warehouse-link-not_linked-20260907-1010.xlsx")
output = Path(r"./data/warehouse-link-20260907-1010.json")
workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
sheet = workbook[workbook.sheetnames[0]]
rows = list(sheet.iter_rows(values_only=True))
headers = [str(value or "").strip() for value in rows[0]]
records = [dict(zip(headers, row)) for row in rows[1:] if any(value not in (None, "") for value in row)]
output.write_text(json.dumps(records, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
print(json.dumps({"source": str(source), "sheet": sheet.title, "rows": len(records), "headers": headers}, ensure_ascii=True))
